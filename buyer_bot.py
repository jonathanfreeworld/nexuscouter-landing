from io import BytesIO
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
import openai
import streamlit as st

# 페이지 설정
st.set_page_config(
    page_title="nexuscouter | AI 어드바이저",
    page_icon="🌐",
    layout="wide",
)

# Streamlit Secrets에서 설정값 안전하게 불러오기
try:
    API_BASE_URL = st.secrets["API_BASE_URL"]
    API_KEY = st.secrets["API_KEY"]
    MODEL_NAME = st.secrets.get("MODEL_NAME", "gpt-6-luna")
except Exception:
    API_BASE_URL = ""
    API_KEY = ""
    MODEL_NAME = "gpt-6-luna"


# 텍스트를 한글 지원 폰트가 적용된 깔끔한 워드 파일로 변환하는 함수
def create_word_document(title: str, content: str) -> bytes:
    doc = docx.Document()

    # 1. 문서 전체의 기본 폰트를 '맑은 고딕'으로 설정 (한글 깨짐 방지)
    style = doc.styles["Normal"]
    font = style.font
    font.name = "맑은 고딕"
    font._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    font.size = Pt(10.5)

    # 문서 제목 추가
    h = doc.add_heading(title, level=1)
    for run in h.runs:
        run.font.name = "맑은 고딕"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")

    doc.add_paragraph("--------------------------------------------------")

    # 2. 마크다운 기호(#, **)를 일부 정리하면서 깔끔하게 워드 단락으로 추가
    for line in content.split("\n"):
        stripped = line.strip()
        if not stripped:
            doc.add_paragraph("")  # 빈 줄 추가
            continue

        # 마크다운 헤더(#) 처리
        if stripped.startswith("### "):
            p = doc.add_paragraph()
            run = p.add_run(stripped.replace("### ", ""))
            run.bold = True
            run.font.size = Pt(13)
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            run = p.add_run(stripped.replace("## ", ""))
            run.bold = True
            run.font.size = Pt(14)
        elif stripped.startswith("# "):
            p = doc.add_paragraph()
            run = p.add_run(stripped.replace("# ", ""))
            run.bold = True
            run.font.size = Pt(15)
        else:
            # 일반 텍스트 내의 ** (볼드) 기호 제거 혹은 깔끔하게 출력
            clean_line = stripped.replace("**", "")
            doc.add_paragraph(clean_line)

    # 메모리 스트림에 저장
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


# 상단 탭 메뉴 구성
tab1, tab2 = st.tabs(
    ["🎯 바이어 타겟팅 & 전략 분석", "✉️ 맞춤형 영문 콜드메일 생성"]
)

# 1. 바이어 타겟팅 & 전략 분석 탭
with tab1:
    st.markdown("### 🎯 해외 바이어 페르소나 및 시장 분석")
    st.write(
        "해외 진출을 위한 타겟 시장과 바이어 조건을 설정하여 맞춤형 전략을 확인하세요."
    )

    col1, col2 = st.columns(2)

    with col1:
        product_name = st.text_input(
            "수출 제품명 (또는 카테고리)",
            placeholder="예: 프리미엄 기능성 화장품 / 제약 원료",
            key="t1_product",
        )
        target_type = st.selectbox(
            "타겟 바이어 유형",
            [
                "현지 총판 디스트리뷰터 (Distributor)",
                "수입사 / 도매상 (Importer / Wholesaler)",
                "현지 온·오프라인 리테일러",
                "기타 파트너",
            ],
            key="t1_type",
        )

    with col2:
        target_country = st.text_input(
            "타겟 국가",
            placeholder="예: 인도, 베트남, 독일 등",
            key="t1_country",
        )
        product_strength = st.text_input(
            "제품 차별점 / 강점",
            placeholder="예: 칼라테라피, 미백, 주름개선, 식물성 추출원료",
            key="t1_strength",
        )

    st.write("")
    if st.button("🚀 바이어 발굴 전략 분석 실행", key="btn_strategy"):
        if product_name and target_country:
            if not API_KEY or not API_BASE_URL:
                st.error(
                    "⚠️ Streamlit Secrets에 API 설정(API_BASE_URL, API_KEY)이 누락되었습니다."
                )
            else:
                st.info(f"[{target_country}] 진출 전략을 실시간 생성 중입니다...")
                try:
                    client = openai.OpenAI(base_url=API_BASE_URL, api_key=API_KEY)

                    prompt = (
                        f"수출 제품: {product_name} / 타겟 국가: {target_country} / "
                        f"바이어 유형: {target_type} / 강점: {product_strength}\n위 조건을 "
                        "바탕으로 핵심 바이어 타겟팅 및 수출 진출 전략을 마크다운으로 작성해 주세요."
                    )

                    stream = client.chat.completions.create(
                        model=MODEL_NAME,
                        messages=[{"role": "user", "content": prompt}],
                        stream=True,
                    )

                    st.markdown("---")

                    full_response = ""
                    placeholder = st.empty()

                    for chunk in stream:
                        if (
                            chunk.choices
                            and chunk.choices[0].delta
                            and chunk.choices[0].delta.content
                        ):
                            content_chunk = chunk.choices[0].delta.content
                            full_response += content_chunk
                            placeholder.markdown(full_response)

                    st.session_state["strategy_result"] = full_response

                except Exception as e:
                    st.error(f"API 호출 중 오류가 발생했습니다: {e}")
        else:
            st.warning("수출 제품명과 타겟 국가를 입력해 주세요.")

    # 워드 다운로드 버튼
    if "strategy_result" in st.session_state and st.session_state["strategy_result"]:
        docx_data = create_word_document(
            f"[{target_country}] 해외 바이어 발굴 및 진출 전략 보고서",
            st.session_state["strategy_result"],
        )
        st.download_button(
            label="📥 해외 진출 전략 보고서 워드 다운로드 (.docx)",
            data=docx_data,
            file_name=f"export_strategy_{target_country or 'report'}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            key="download_strategy_docx",
        )

# 2. 맞춤형 영문 콜드메일 생성 탭
with tab2:
    st.markdown("### ✉️ 바이오 맞춤형 영문 콜드메일(Cold Mail) 자동 생성")
    st.write(
        "발견한 바이어에게 보낼 매력적인 첫인상 영문 메일 초안을 자동으로 작성합니다."
    )

    col3, col4 = st.columns(2)

    with col3:
        company_name = st.text_input(
            "우리회사 명칭", placeholder="예: KJ H&B", key="t2_company"
        )
        buyer_name = st.text_input(
            "바이어 회사명",
            placeholder="예: Apex Pharma India",
            key="t2_buyer",
        )

    with col4:
        item_content = st.text_input(
            "제안 품목 핵심 내용",
            placeholder="예: 기능성 화장품 및 식물성 추출원료",
            key="t2_item",
        )
        tone_manner = st.selectbox(
            "톤앤매너",
            [
                "전문적이고 격식 있는 (Formal)",
                "친근하고 설득력 있는 (Friendly & Persuasive)",
                "간결하고 명확한 (Concise & Direct)",
            ],
            key="t2_tone",
        )

    st.write("")
    if st.button("✨ 영문 제안 메일 생성하기", key="btn_mail"):
        if company_name and buyer_name and item_content:
            if not API_KEY or not API_BASE_URL:
                st.error(
                    "⚠️ Streamlit Secrets에 API 설정(API_BASE_URL, API_KEY)이 누락되었습니다."
                )
            else:
                st.info("영문 콜드메일 초안을 실시간 작성 중입니다...")
                try:
                    client = openai.OpenAI(base_url=API_BASE_URL, api_key=API_KEY)
                    mail_prompt = (
                        f"보내는 회사: {company_name} / 수신 바이어: {buyer_name} / 제안 "
                        f"품목: {item_content} / 톤앤매너: {tone_manner}\n위 정보를 "
                        "바탕으로 영문 콜드메일 초안(Subject 및 본문)을 작성해 주세요."
                    )

                    stream = client.chat.completions.create(
                        model=MODEL_NAME,
                        messages=[{"role": "user", "content": mail_prompt}],
                        stream=True,
                    )

                    st.markdown("---")

                    full_mail = ""
                    mail_placeholder = st.empty()

                    for chunk in stream:
                        if (
                            chunk.choices
                            and chunk.choices[0].delta
                            and chunk.choices[0].delta.content
                        ):
                            content_chunk = chunk.choices[0].delta.content
                            full_mail += content_chunk
                            mail_placeholder.markdown(full_mail)

                    st.session_state["mail_result"] = full_mail

                except Exception as e:
                    st.error(f"API 호출 중 오류가 발생했습니다: {e}")
        else:
            st.warning("필수 입력 항목을 모두 작성해 주세요.")

    # 워드 다운로드 버튼
    if "mail_result" in st.session_state and st.session_state["mail_result"]:
        docx_mail_data = create_word_document(
            f"[{buyer_name}] 영문 콜드메일 제안서",
            st.session_state["mail_result"],
        )
        st.download_button(
            label="📥 영문 콜드메일 초안 워드 다운로드 (.docx)",
            data=docx_mail_data,
            file_name=f"cold_mail_{buyer_name or 'buyer'}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            key="download_mail_docx",
        )
